from __future__ import annotations

import sqlite3
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app import export


def _template(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BOQ"
    sheet["B1"] = "Nội dung"
    sheet["E1"] = "Khối lượng"
    sheet["F1"] = "Đơn giá vật tư"
    sheet["G1"] = "Đơn giá nhân công"
    sheet["H1"] = "Thành tiền vật tư"
    sheet["I1"] = "Thành tiền nhân công"
    sheet["J1"] = "Tổng"
    sheet["B2"] = "Cáp CXV"
    sheet["E2"] = 2
    sheet["H2"] = "=E2*F2"
    sheet["I2"] = "=E2*G2"
    sheet["J2"] = "=SUM(H2,I2)"
    workbook.save(path)


def test_export_only_updates_unit_price_columns(monkeypatch, tmp_path: Path) -> None:
    source_path = tmp_path / "input.xlsx"
    _template(source_path)
    output_dir = tmp_path / "exports"
    output_dir.mkdir()

    conn = sqlite3.connect(":memory:")
    conn.row_factory = sqlite3.Row
    conn.executescript(
        """
        CREATE TABLE projects(id INTEGER PRIMARY KEY, source_file_id INTEGER);
        CREATE TABLE source_files(id INTEGER PRIMARY KEY, storage_key TEXT);
        CREATE TABLE source_sheets(
            id INTEGER PRIMARY KEY, sheet_name TEXT, mapping_json TEXT
        );
        CREATE TABLE source_rows(id INTEGER PRIMARY KEY, row_no INTEGER);
        INSERT INTO projects VALUES (1, 1);
        INSERT INTO source_rows VALUES (1, 2);
        """
    )
    conn.execute(
        "INSERT INTO source_files VALUES (1, ?)",
        (str(source_path),),
    )
    conn.execute(
        "INSERT INTO source_sheets VALUES (1, 'BOQ', ?)",
        (
            '{"description":1,"quantity":4,"material_price":5,'
            '"labor_price":6,"total":7,"amount":9}',
        ),
    )

    class Session:
        def __enter__(self):
            return conn

        def __exit__(self, exc_type, exc, traceback):
            return False

    monkeypatch.setattr(export, "EXPORT_DIR", output_dir)
    monkeypatch.setattr(export, "ensure_directories", lambda: None)
    monkeypatch.setattr(export, "db_session", lambda: Session())
    monkeypatch.setattr(
        export,
        "get_project_result",
        lambda project_id, run_id: {
            "items": [
                {
                    "id": 1,
                    "source_sheet_id": 1,
                    "source_row_id": 1,
                    "raw_description": "Cáp CXV",
                    "unit": "m",
                    "quantity": 2,
                    "material_price": 100,
                    "labor_price": 20,
                    "material_total": 200,
                    "labor_total": 40,
                    "status": "AUTO_APPROVED",
                }
            ]
        },
    )

    output_path = export.export_project(1)
    workbook = load_workbook(output_path, data_only=False)
    sheet = workbook["BOQ"]

    assert sheet["F2"].value == 100
    assert sheet["G2"].value == 20
    assert sheet["H2"].value == "=E2*F2"
    assert sheet["I2"].value == "=E2*G2"
    assert sheet["J2"].value == "=SUM(H2,I2)"
