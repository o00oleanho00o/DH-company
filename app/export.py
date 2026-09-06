from __future__ import annotations

import json
import os
import re
import shutil
import tempfile
import zipfile
from xml.etree import ElementTree as ET
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from .config import EXPORT_DIR, ensure_directories
from .db import db_session, dumps, loads, utc_now
from .pricing import get_project_result


def _column(mapping: dict[str, Any], key: str) -> int | None:
    value = mapping.get(key)
    try:
        return int(value) + 1 if value is not None else None
    except (TypeError, ValueError):
        return None


def _write_cell(ws, row_no: int, col_no: int | None, value: Any) -> None:
    if col_no and row_no > 0:
        ws.cell(row=row_no, column=col_no).value = value


def _audit_rows(result: dict[str, Any]) -> list[list[Any]]:
    rows = [[
        "BOQ item", "Mô tả gốc", "Đơn vị", "Khối lượng",
        "Giá vật tư", "Nguồn vật tư", "Confidence VT",
        "Giá nhân công", "Nguồn nhân công", "Confidence NC",
        "Trạng thái", "Rủi ro", "Giải thích",
    ]]
    for item in result.get("items", []):
        ms = item.get("material_source") or {}
        ls = item.get("labor_source") or {}
        rows.append([
            item.get("id"),
            item.get("raw_description"),
            item.get("unit"),
            item.get("quantity"),
            item.get("material_price"),
            _source_label(ms),
            item.get("material_confidence"),
            item.get("labor_price"),
            _source_label(ls),
            item.get("labor_confidence"),
            item.get("status"),
            item.get("risk"),
            item.get("explanation"),
        ])
    return rows


def _source_label(source: dict[str, Any]) -> str:
    if not source:
        return ""
    parts = []
    if source.get("supplier"):
        parts.append(str(source["supplier"]))
    if source.get("filename"):
        parts.append(str(source["filename"]))
    if source.get("sheet_name"):
        parts.append(f"sheet:{source['sheet_name']}")
    if source.get("row_no"):
        parts.append(f"row:{source['row_no']}")
    if source.get("effective_date"):
        parts.append(f"date:{source['effective_date']}")
    if source.get("rate"):
        parts.append(f"rate:{source['rate']}")
    return " | ".join(parts)


def _style_audit(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {
        "A": 12, "B": 62, "C": 12, "D": 14, "E": 16, "F": 52,
        "G": 16, "H": 16, "I": 52, "J": 16, "K": 18, "L": 12, "M": 52,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


_XLSX_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_XLSX_NS = {"x": _XLSX_MAIN_NS}


def _formula_cache_map(path: Path) -> dict[str, dict[str, tuple[str | None, str | None]]]:
    """Read formula + cached-value pairs directly from the source XML.

    openpyxl intentionally does not preserve cached formula results when a
    workbook is rewritten. Reading the original package lets us restore those
    values after adding the audit sheet and deterministic prices.
    """

    result: dict[str, dict[str, tuple[str | None, str | None]]] = {}
    try:
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                if not (name.startswith("xl/worksheets/") and name.endswith(".xml")):
                    continue
                root = ET.fromstring(archive.read(name))
                cells: dict[str, tuple[str | None, str | None]] = {}
                for cell in root.findall(".//x:c", _XLSX_NS):
                    coordinate = cell.attrib.get("r")
                    formula = cell.find("x:f", _XLSX_NS)
                    cached = cell.find("x:v", _XLSX_NS)
                    if coordinate and formula is not None and cached is not None:
                        cells[coordinate] = (
                            formula.text,
                            cached.text,
                        )
                if cells:
                    result[name] = cells
    except (OSError, ET.ParseError, zipfile.BadZipFile):
        return {}
    return result


def _restore_formula_caches(source_path: Path, output_path: Path) -> None:
    """Restore cached values for formula cells that survived export edits."""

    cache_map = _formula_cache_map(source_path)
    if not cache_map:
        return
    fd, temp_name = tempfile.mkstemp(
        prefix=f"{output_path.stem}-cache-",
        suffix=".xlsx",
        dir=str(output_path.parent),
    )
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        with zipfile.ZipFile(output_path, "r") as source_zip, zipfile.ZipFile(
            temp_path, "w", compression=zipfile.ZIP_DEFLATED
        ) as target_zip:
            for info in source_zip.infolist():
                payload = source_zip.read(info.filename)
                cells = cache_map.get(info.filename)
                if cells:
                    try:
                        root = ET.fromstring(payload)
                    except ET.ParseError:
                        root = None
                    if root is not None:
                        for cell in root.findall(".//x:c", _XLSX_NS):
                            coordinate = cell.attrib.get("r")
                            cached = cells.get(coordinate or "")
                            formula = cell.find("x:f", _XLSX_NS)
                            if cached and formula is not None:
                                source_formula, source_value = cached
                                # Only restore a cache when the formula itself
                                # remains unchanged. Exported total/amount
                                # cells may have intentionally become values.
                                if source_formula == formula.text and source_value is not None:
                                    target = cell.find("x:v", _XLSX_NS)
                                    if target is None:
                                        target = ET.SubElement(
                                            cell,
                                            f"{{{_XLSX_MAIN_NS}}}v",
                                        )
                                    target.text = source_value
                        payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                target_zip.writestr(info, payload)
        os.replace(temp_path, output_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


_CELL_REF_RE = re.compile(r"(?<![A-Z0-9_])(?:'[^']+'!)?\$?([A-Z]{1,3})\$?(\d+)")
_SUM_RE = re.compile(
    r"SUM\(\s*(?:'[^']+'!)?\$?([A-Z]{1,3})\$?(\d+)"
    r"\s*:\s*(?:'[^']+'!)?\$?([A-Z]{1,3})\$?(\d+)\s*\)",
    re.IGNORECASE,
)
_SUM_ARGS_RE = re.compile(r"SUM\(([^()]*)\)", re.IGNORECASE)
_ROUND_RE = re.compile(r"ROUND\(([^(),]+),\s*(-?\d+)\)", re.IGNORECASE)


def _recalculate_formula_caches(output_path: Path) -> None:
    """Refresh caches for simple formulas after deterministic workbook edits."""

    try:
        wb = load_workbook(output_path, data_only=False)
    except (OSError, ValueError):
        return
    # Rewriting worksheet XML to inject cached values is only needed for the
    # compact single-BOQ template whose grand totals otherwise retain a zero
    # cache. Large multi-sheet workbooks contain shared formulas, extension
    # metadata and cross-sheet calculations that Excel validates strictly.
    # Leave those formulas to Excel's native recalculation engine instead of
    # serializing every worksheet through ElementTree.
    business_sheets = [
        ws for ws in wb.worksheets if ws.title != "AI Audit"
    ]
    if len(business_sheets) != 1:
        return

    memo: dict[tuple[int, str], float | None] = {}
    active: set[tuple[int, str]] = set()

    def formula_value(sheet_index: int, formula: str) -> float | None:
        expression = formula.lstrip("=").strip()
        if expression.startswith("+"):
            expression = expression[1:].strip()

        def replace_sum(match: re.Match[str]) -> str:
            start_col, start_row, end_col, end_row = match.groups()
            ws = wb.worksheets[sheet_index]
            total = 0.0
            for row in ws.iter_rows(
                min_row=int(start_row),
                max_row=int(end_row),
                min_col=ws[f"{start_col}1"].column,
                max_col=ws[f"{end_col}1"].column,
            ):
                for cell in row:
                    value = cell_value(sheet_index, cell.coordinate)
                    if value is not None:
                        total += value
            return str(total)

        expression = _SUM_RE.sub(replace_sum, expression)

        def replace_sum_args(match: re.Match[str]) -> str:
            values = []
            for token in match.group(1).split(","):
                token = token.strip()
                try:
                    values.append(float(token))
                except ValueError:
                    return match.group(0)
            return str(sum(values))

        expression = _SUM_ARGS_RE.sub(replace_sum_args, expression)

        def replace_ref(match: re.Match[str]) -> str:
            token = match.group(0)
            if "!" in token:
                return token
            col, row = match.groups()
            value = cell_value(sheet_index, f"{col}{row}")
            return str(value) if value is not None else token

        expression = _CELL_REF_RE.sub(replace_ref, expression)
        def replace_round(match: re.Match[str]) -> str:
            try:
                return str(round(float(match.group(1)), int(match.group(2))))
            except (TypeError, ValueError):
                return match.group(0)

        expression = _ROUND_RE.sub(replace_round, expression)
        if _CELL_REF_RE.search(expression):
            return None
        if not re.fullmatch(r"[0-9eE+\-*/().\s]+", expression):
            return None
        try:
            value = float(eval(expression, {"__builtins__": {}}, {}))
        except (ArithmeticError, SyntaxError, ValueError, TypeError):
            return None
        return value if value == value and abs(value) != float("inf") else None

    def cell_value(sheet_index: int, coordinate: str) -> float | None:
        key = (sheet_index, coordinate)
        if key in memo:
            return memo[key]
        if key in active:
            return None
        active.add(key)
        try:
            value = wb.worksheets[sheet_index][coordinate].value
            if isinstance(value, bool):
                result = float(value)
            elif isinstance(value, (int, float)):
                result = float(value)
            elif isinstance(value, str) and value.startswith("="):
                result = formula_value(sheet_index, value)
            else:
                result = None
        finally:
            active.discard(key)
        memo[key] = result
        return result

    updates: dict[str, dict[str, float]] = {}
    for sheet_index, ws in enumerate(wb.worksheets):
        xml_name = f"xl/worksheets/sheet{sheet_index + 1}.xml"
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    value = cell_value(sheet_index, cell.coordinate)
                    if value is not None:
                        updates.setdefault(xml_name, {})[cell.coordinate] = value
    if not updates:
        return

    fd, temp_name = tempfile.mkstemp(
        prefix=f"{output_path.stem}-recalc-",
        suffix=".xlsx",
        dir=str(output_path.parent),
    )
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        with zipfile.ZipFile(output_path, "r") as source_zip, zipfile.ZipFile(
            temp_path, "w", compression=zipfile.ZIP_DEFLATED
        ) as target_zip:
            for info in source_zip.infolist():
                payload = source_zip.read(info.filename)
                sheet_updates = updates.get(info.filename)
                if sheet_updates:
                    try:
                        root = ET.fromstring(payload)
                    except ET.ParseError:
                        root = None
                    if root is not None:
                        for cell in root.findall(".//x:c", _XLSX_NS):
                            value = sheet_updates.get(cell.attrib.get("r") or "")
                            if value is None:
                                continue
                            cached = cell.find("x:v", _XLSX_NS)
                            if cached is None:
                                cached = ET.SubElement(cell, f"{{{_XLSX_MAIN_NS}}}v")
                            cached.text = format(value, ".15g")
                        payload = ET.tostring(
                            root, encoding="utf-8", xml_declaration=True
                        )
                target_zip.writestr(info, payload)
        os.replace(temp_path, output_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()


def export_project(project_id: int, run_id: int | None = None) -> Path:
    """Create a usable XLSX result and append a provenance-rich AI Audit sheet."""

    ensure_directories()
    result = get_project_result(project_id, run_id)
    with db_session() as conn:
        project = conn.execute("SELECT * FROM projects WHERE id=?", (project_id,)).fetchone()
        if not project:
            raise KeyError(project_id)
        source = conn.execute(
            "SELECT * FROM source_files WHERE id=?", (project["source_file_id"],)
        ).fetchone()
        source_path = Path(source["storage_key"]) if source and source["storage_key"] else None
        output_path = EXPORT_DIR / f"quotation-{project_id}-{run_id or 'latest'}.xlsx"

        preserve_formula_source = (
            source_path
            if source_path and source_path.exists() and source_path.suffix.lower() == ".xlsx"
            else None
        )
        if preserve_formula_source:
            shutil.copy2(source_path, output_path)
            wb = load_workbook(output_path)
            # Source sheet metadata stores the deterministic mapping selected at
            # ingest time, so export does not depend on fixed columns.
            for item in result.get("items", []):
                if not item.get("source_sheet_id") or not item.get("source_row_id"):
                    continue
                ss = conn.execute(
                    "SELECT * FROM source_sheets WHERE id=?", (item["source_sheet_id"],)
                ).fetchone()
                sr = conn.execute(
                    "SELECT row_no FROM source_rows WHERE id=?", (item["source_row_id"],)
                ).fetchone()
                if not ss or not sr or ss["sheet_name"] not in wb.sheetnames:
                    continue
                ws = wb[ss["sheet_name"]]
                mapping = loads(ss["mapping_json"], {})
                row_no = int(sr["row_no"])
                _write_cell(ws, row_no, _column(mapping, "material_price"), item.get("material_price"))
                _write_cell(ws, row_no, _column(mapping, "labor_price"), item.get("labor_price"))
                # Generic total columns are only written when the mapping is
                # unambiguous; formulas are replaced by deterministic values.
                total = None
                if item.get("material_total") is not None or item.get("labor_total") is not None:
                    total = (item.get("material_total") or 0) + (item.get("labor_total") or 0)
                _write_cell(ws, row_no, _column(mapping, "total"), total)
                _write_cell(ws, row_no, _column(mapping, "amount"), total)
                # In the supplied BOQ template the column immediately before
                # the amount/total column is the combined unit price (I),
                # while G/H hold the material/labor split. Rebuild that
                # deterministic value when the mapping identifies such a
                # column, so the exported workbook remains usable in Excel
                # even when the input holdout intentionally omitted prices.
                total_col = _column(mapping, "total")
                material_col = _column(mapping, "material_price")
                labor_col = _column(mapping, "labor_price")
                if total_col and total_col > 1:
                    combined_col = total_col - 1
                    if combined_col not in {
                        value for value in (material_col, labor_col, _column(mapping, "quantity"))
                        if value
                    }:
                        combined_unit = None
                        if item.get("material_price") is not None or item.get("labor_price") is not None:
                            combined_unit = (item.get("material_price") or 0) + (
                                item.get("labor_price") or 0
                            )
                        _write_cell(ws, row_no, combined_col, combined_unit)
        else:
            # Legacy .xls cannot be safely edited with openpyxl. Produce a
            # normalized, fully usable workbook instead of silently failing.
            wb = Workbook()
            ws = wb.active
            ws.title = "BOQ kết quả"
            headers = [
                "STT", "Mô tả", "Mã", "Đơn vị", "Khối lượng",
                "Đơn giá vật tư", "Đơn giá nhân công", "Tổng vật tư",
                "Tổng nhân công", "Tổng cộng", "Trạng thái", "Nguồn",
            ]
            ws.append(headers)
            for i, item in enumerate(result.get("items", []), 1):
                total = (item.get("material_total") or 0) + (item.get("labor_total") or 0)
                ws.append([
                    i, item.get("raw_description"), item.get("product_code"),
                    item.get("unit"), item.get("quantity"), item.get("material_price"),
                    item.get("labor_price"), item.get("material_total"),
                    item.get("labor_total"), total, item.get("status"),
                    _source_label(item.get("material_source") or {}) or _source_label(item.get("labor_source") or {}),
                ])
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            for col, width in {"A": 8, "B": 70, "C": 20, "D": 12, "E": 14, "F": 18, "G": 18, "H": 18, "I": 18, "J": 18, "K": 18, "L": 60}.items():
                ws.column_dimensions[col].width = width

        if "AI Audit" in wb.sheetnames:
            del wb["AI Audit"]
        audit = wb.create_sheet("AI Audit")
        for row in _audit_rows(result):
            audit.append(row)
        _style_audit(audit)
        wb.save(output_path)
        if preserve_formula_source:
            _restore_formula_caches(preserve_formula_source, output_path)
            _recalculate_formula_caches(output_path)
        return output_path
